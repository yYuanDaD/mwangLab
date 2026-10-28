"""Paper screening: keep only papers whose deposited data can support a proper analysis.

Goal: filter at the input instead of at review time. A paper enters the analysis run only
when its full text is readable, it deposits its OWN data, the deposit is a supported data
type, the data files needed for that route exist, and the sample design has at least two
groups of three. Every paper gets one row with the verdict, the reason, and the evidence
(block IDs, file names, group sizes), so the loss at each layer can be counted.

Layers (each cheap enough to run on a few hundred papers; nothing here calls an LLM):

    1. text      full text via Europe PMC / NCBI PMC JATS, else the PDF (tools.pdf_markdown);
                 the title or abstract must name the topic (exercise by default), and the
                 abstract, methods or results must describe the paper's own exercise
                 groups or protocol, not just mention exercise; own vs.
                 cited accessions from the words around each mention, then the section
    2. deposit   GEO series summary (E-utilities esummary): data type, organism, sample
                 count, platform; organism (human, mouse, rat by default); non-GEO repositories (PRIDE, ArrayExpress, SRA-only,
                 GSA) are recorded and routed, not analysed
    3. files     GEO supplementary listing: raw counts vs. TPM/FPKM (limma) vs. per-sample
                 tar vs. nothing (bulk RNA-seq only); values and column names decide over
                 file names; a _RAW.tar is judged from its per-sample files (listed in the
                 series matrix), sniffing one or two
    4. design    GEO sample characteristics (series matrix): two compared groups of >= 3
                 (small extra groups are flagged, not fatal), sex composition and the arms
                 per sex; uses metadata_structural's deterministic design-column picker

GEO's own PMID -> series links (elink) add the series GEO links to the paper first, so a
paper whose text omits its accession, or whose full text was not fetched, is not lost.

Verdicts: ``analyzable`` · ``needs_inspection`` (a layer could not decide, e.g. only a
``_RAW.tar``) · ``excluded`` with a reason code.

Network access goes through one ``Http`` object with an on-disk cache, so reruns are
free and tests inject canned responses.

CLI (repository root):

    python -m tools.paper_screen collect "exercise skeletal muscle" "endurance training" \
        --sources s2,pubmed,geo --out output/screen_run --max 1000
    python -m tools.paper_screen screen output/screen_run/papers.csv --out output/screen_run
    python -m tools.paper_screen summary output/screen_run/screen.csv
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import io
import json
import os
import re
import sys
import time
from collections import Counter
from dataclasses import dataclass, field

EUTILS = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils"
GEO_SUPPL = "https://ftp.ncbi.nlm.nih.gov/geo/series/{stub}/{gse}/suppl/"
GEO_MATRIX = "https://ftp.ncbi.nlm.nih.gov/geo/series/{stub}/{gse}/matrix/"
S2_SEARCH = "https://api.semanticscholar.org/graph/v1/paper/search"
_UA = {"User-Agent": "mwangLab-paper-screen/1.0 (bioinformatics research)"}

# Section classes where a GSE mention usually means the paper's own deposit.
OWN_SECTIONS = {"methods", "data_availability", "supplementary", "other"}
CITED_SECTIONS = {"introduction", "discussion"}

_REPO_PATTERNS = {
    "geo": re.compile(r"\bGSE\d{3,}\b"),
    "pride": re.compile(r"\bPXD\d{6}\b"),
    "arrayexpress": re.compile(r"\bE-[A-Z]{4}-\d+\b"),
    "sra": re.compile(r"\b(?:PRJNA\d{3,}|PRJEB\d{3,}|SRP\d{4,})\b"),
    "gsa": re.compile(r"\b(?:CRA\d{6}|HRA\d{6}|PRJCA\d{6})\b"),
    "dbgap": re.compile(r"\bphs\d{6}\b"),
}
_ON_REQUEST = re.compile(r"(available|provided|obtained)\s+(from the corresponding author\s+)?"
                         r"(up)?on\s+(reasonable\s+)?request", re.I)

_GDSTYPE = (
    ("expression profiling by high throughput sequencing", "rnaseq"),
    ("expression profiling by array", "microarray"),
    ("methylation profiling", "methylation"),
    ("non-coding rna profiling by high throughput sequencing", "ncrna_seq"),
    ("non-coding rna profiling by array", "ncrna_array"),
    ("genome binding/occupancy", "chromatin"),
    ("genome variation", "genome_variation"),
    ("protein profiling", "proteomics"),
)
_SINGLE_CELL = re.compile(r"single[- ]?(cell|nucle(us|i))|scRNA|snRNA|10x genomics|chromium|drop-?seq", re.I)

# Supplementary file names.
_RAW_COUNT_NAME = re.compile(
    r"(raw[._ -]?counts?|read[._ -]?counts?|counts?[._ -]?(matrix|table)|featurecounts?|htseq|"
    # "counts" as a word of its own, where "_" also separates words ("X_Full_Counts_v2.xlsx").
    r"gene[._ -]?counts?|(?<![a-z])counts?(?![a-z])|expected[._ -]?count|umi[._ -]?counts?|"
    # "SI_allcounts", "CounTable", and "raw" as a word of its own ("silver_hmt_raw.csv.gz").
    r"allcounts?|coun?t?table|(?<![a-z])raw(?![a-z]))", re.I)
# "norm" joined to the next word too: "DESeq2_normcounts.xlsx", "Cortex_EV_NormCounts.xlsx".
_NORMALISED_NAME = re.compile(r"(fpkm|rpkm|tpm|cpm|(?<![a-z])tmm(?![a-z])|normali[sz]ed|norm(?:[._-]|counts?|expr|data)|vst|rlog|log2|"
                              r"abundance)", re.I)
# Differential-expression result tables: they cannot feed a re-analysis.
_DE_NAME = re.compile(r"((?<![a-z])(degs?|dges?)(?![a-z])|diff[._ -]?expr|differential|"
                      r"(?<![a-z])de[._ -](summary|results?|tables?|genes?)|log2?fc|fold[._ -]?change)", re.I)
_MATRIX_EXT = re.compile(r"\.(csv|tsv|txt|xlsx?|tab)(\.gz)?$", re.I)


# ---------------------------------------------------------------------------
# HTTP with an on-disk cache
# ---------------------------------------------------------------------------

class Http:
    """GET with retries, NCBI rate limiting, and a response cache keyed by URL+params."""

    def __init__(self, cache_dir: str = "", api_key: str = "", min_interval: float = 0.0):
        self.cache_dir = cache_dir
        if cache_dir:
            os.makedirs(cache_dir, exist_ok=True)
        self.api_key = api_key or os.getenv("NCBI_API_KEY", "")
        self.min_interval = min_interval or (0.11 if self.api_key else 0.35)
        self._last = 0.0

    def _key(self, url: str, params: dict | None) -> str:
        raw = url + "?" + json.dumps(params or {}, sort_keys=True)
        return os.path.join(self.cache_dir, hashlib.sha1(raw.encode()).hexdigest()[:20])

    def get(self, url: str, params: dict | None = None, binary: bool = False, timeout: int = 60,
            max_bytes: int = 0):
        """Body (str, or bytes when binary) or None on 404 / repeated failure.

        ``max_bytes``: stop after this many bytes (the head of a large matrix file)."""
        params = dict(params or {})
        if url.startswith(EUTILS) and self.api_key:
            params["api_key"] = self.api_key
        key_params = {k: v for k, v in params.items() if k != "api_key"}
        if max_bytes:
            key_params["__max_bytes__"] = max_bytes
        path = self._key(url, key_params) if self.cache_dir else ""
        if path and os.path.exists(path):
            with open(path, "rb") as f:
                data = f.read()
            if data == b"\x00404":
                return None
            return data if binary else data.decode("utf-8", errors="replace")
        import requests
        headers = dict(_UA)
        s2 = url.startswith(S2_SEARCH)
        if s2 and os.getenv("S2_API_KEY"):
            headers["x-api-key"] = os.environ["S2_API_KEY"]
        data = None
        for attempt in range(4):
            wait = self.min_interval - (time.time() - self._last)
            if wait > 0:
                time.sleep(wait)
            self._last = time.time()
            try:
                r = requests.get(url, params=params, headers=headers, timeout=timeout, stream=bool(max_bytes))
            except Exception:
                time.sleep(2 ** attempt)
                continue
            if r.status_code in (401, 403, 404, 410):
                # Not there, or a publisher blocking automated download: unavailable,
                # not an error (a 403 PDF means "no full text", not "inspect by hand").
                data = b"\x00404"
                break
            if r.status_code == 500 and "europepmc" in url and attempt >= 1:
                # Europe PMC answers 500 (not 404) for full text it may not serve; cache
                # it so a rerun does not spend the full retry backoff on every such paper.
                data = b"\x00404"
                break
            if r.status_code == 429 or r.status_code >= 500:
                # The shared anonymous Semantic Scholar pool throttles for tens of seconds.
                time.sleep(20 * (attempt + 1) if s2 and r.status_code == 429 else 2 ** (attempt + 1))
                continue
            r.raise_for_status()
            if max_bytes:
                data = b""
                for chunk in r.iter_content(65536):
                    data += chunk
                    if len(data) >= max_bytes:
                        break
                r.close()
            else:
                data = r.content
            break
        if data is None:
            return None
        if path:
            with open(path, "wb") as f:
                f.write(data)
        if data == b"\x00404":
            return None
        return data if binary else data.decode("utf-8", errors="replace")


# ---------------------------------------------------------------------------
# Layer 1: own vs cited accessions from structured full text
# ---------------------------------------------------------------------------

def deposits_from_doc(doc: dict) -> dict:
    """Accessions the paper deposits, from a structured paper (tools.paper_structure).

    A GEO series counts as the paper's own when it is mentioned in Methods, the data
    availability statement, supplementary material or front matter. One that appears only
    in the Introduction or Discussion is a citation. Returns
    ``{"own": {repo: [ids]}, "cited": [GSE...], "evidence": {id: [block ids]},
    "on_request": bool}``.
    """
    by_id = {b["id"]: b for b in doc.get("blocks", [])}
    where: dict[str, set] = {}
    evidence: dict[str, list] = {}
    on_request = False
    for b in doc.get("blocks", []):
        cls = b.get("section_class", "other")
        for repo, pat in _REPO_PATTERNS.items():
            for acc in dict.fromkeys(pat.findall(b["text"])):
                where.setdefault(acc, set()).add(cls)
                evidence.setdefault(acc, []).append(b["id"])
        if cls in ("data_availability", "methods") and _ON_REQUEST.search(b["text"]):
            on_request = True
    own: dict[str, list] = {}
    cited = []
    said_own = set()
    for acc, classes in where.items():
        repo = next(r for r, p in _REPO_PATTERNS.items() if p.fullmatch(acc))
        # The words around each mention decide first: "deposited / generated in this
        # study" is the paper's own; "downloaded from / publicly available / re-analysed"
        # is reuse, even in Methods or the data-availability statement. Only when no
        # mention carries either phrase does the section decide.
        phrases = {_mention_phrase(acc, bid, by_id) for bid in evidence[acc]}
        if "own" in phrases:
            own.setdefault(repo, []).append(acc)
            said_own.add(acc)
        elif "cited" in phrases:
            cited.append(acc)
        elif classes & OWN_SECTIONS or not classes <= CITED_SECTIONS:
            own.setdefault(repo, []).append(acc)
        else:
            cited.append(acc)
    # A meta-analysis lists dozens of series in its tables and Methods with no deposit
    # wording (MetaMEx, PMC6981202: ~50). A paper that deposits that many series says
    # so; past MANY_SERIES, keep only the ones with deposit wording.
    if len(own.get("geo", [])) > MANY_SERIES:
        cited += [a for a in own["geo"] if a not in said_own]
        own["geo"] = [a for a in own["geo"] if a in said_own]
        if not own["geo"]:
            del own["geo"]
    return {"own": own, "cited": cited, "evidence": evidence, "on_request": on_request,
            "said_own": sorted(said_own)}


MANY_SERIES = 8


_CITED_NEAR = re.compile(
    r"(downloaded|obtained|retrieved|accessed|acquired|extracted|compiled)\b|publicly (available|accessible)|"
    r"public (data ?sets?|data|repositor)|previously (published|reported|described|deposited|generated)|"
    r"re-?analy[sz]|meta-?analy[sz]|re-?used?\b|from (a|the|our) (previous|prior|earlier|published)|(previous|prior|earlier) stud|"
    r"(data ?sets?|data|profiles?) (were |was )?(from|of) (the )?(GEO|Gene Expression Omnibus)|"
    r"available (from|in|at|on) (the )?(GEO|Gene Expression Omnibus)|published data|"
    r"(data ?sets?|data|profiles?) (from|of|by|published by|reported by) (the )?[A-Z][\w'-]+ et al", re.I)
_OWN_NEAR = re.compile(
    r"(have been|has been|were|was|are|is|been) (deposited|submitted|uploaded)|\bdeposited (in|at|to|under|with)|"
    r"we (have )?(deposited|submitted|uploaded)|generated (in|for|during) (this|the present|the current) stud|"
    r"(produced|reported|sequenced) (in|for) (this|the present|the current) stud|"
    r"data (from|of|for) this (study|article|paper)|(this|the present|the current) study (are|is) available", re.I)


# "data are publicly available (GSE205676)", "publicly available repositories (GEO: ...)":
# the deposit itself. "analysed publicly available datasets: GSE113165" stays reuse.
_PUBLICLY = re.compile(r"\b(is|are|be|been|was|were)\s+(made\s+|now\s+)?publicly (available|accessible)|"
                       r"publicly (available|accessible)\s+(repositor\w*|databases?|in|at|on|via|through|from the)", re.I)
_LIST_GAP = re.compile(r"[\s,;/()\[\]]*((and|or|&)[\s,;()]*)?", re.I)


def _mention_phrase(acc: str, bid: str, by_id: dict) -> str:
    """'own', 'cited' or '' from the words around one mention of an accession.

    The window stops at the previous accession or sentence so a phrase about another
    dataset in the same paragraph is not borrowed.
    """
    text = by_id[bid]["text"]
    i = text.find(acc)
    if i < 0:
        return ""
    before = text[max(0, i - 300): i]
    accs = list(re.finditer(r"\b(GSE|PXD|E-[A-Z]{4}-|PRJNA|SRP)\w*", before))
    # In a list ("are accessible at GSE59088 and GSE97084") the phrase before the first
    # item covers every item, so step back over accessions joined only by separators.
    end = len(before)
    while accs and _LIST_GAP.fullmatch(before[accs[-1].end():end]):
        end = accs.pop().start()
    cut = max(before.rfind(". ", 0, end), accs[-1].end() if accs else -1)
    window = before[cut + 1:] + text[i: i + 80]
    # "deposited ... GSE294565 (Nolt et al. 2024)": the series came with an earlier paper.
    if _CITE_AFTER.match(text[i + len(acc): i + len(acc) + 60]):
        return "cited"
    if _OWN_NEAR.search(window):
        return "own"
    # "All RNA-seq data are publicly available (GSE205676)" in a data-availability
    # statement describes the paper's own deposit; elsewhere "a publicly available
    # dataset (GSE163356)" is reuse.
    if by_id[bid].get("section_class") == "data_availability":
        window = _PUBLICLY.sub("", window)
    # "the GESTALT cohort^6 (GSE129643)": a reference marker right before the accession
    # credits the dataset to another paper.
    if _CITED_NEAR.search(window) or _REF_BEFORE.search(text[max(0, i - 40): i]):
        return "cited"
    return ""


_CITE_AFTER = re.compile(r"\s*[(\[]\s*[A-Z][A-Za-z'’-]+(?:\s+et\s+al\.?|\s+(?:and|&)\s+[A-Z][A-Za-z'’-]+)?"
                         r",?\s*\(?(?:19|20)\d{2}[a-z]?\s*\)?\s*[)\];]")
_REF_BEFORE = re.compile(r"(\^\d+(?:[,–-]\^?\d+)*|\[\d+(?:[,–-]\d+)*\]|\(\d+(?:[,–-]\d+)*\))\s*[(\[]?\s*$")


# ---------------------------------------------------------------------------
# Layer 2: GEO series summary
# ---------------------------------------------------------------------------

def geo_stub(gse: str) -> str:
    n = gse[3:]
    return f"GSE{n[:-3]}nnn" if len(n) > 3 else "GSEnnn"


def classify_gdstype(gdstype: str, text: str = "") -> str:
    """Data type for routing; several types joined by '+' (e.g. 'rnaseq+chromatin')."""
    low = (gdstype or "").lower()
    kinds = [k for phrase, k in _GDSTYPE if phrase in low]
    if "rnaseq" in kinds and _SINGLE_CELL.search(text or ""):
        kinds[kinds.index("rnaseq")] = "scrnaseq"
    return "+".join(dict.fromkeys(kinds)) or ("other" if low else "unknown")


def geo_summary(gse: str, http: Http) -> dict | None:
    """esummary fields for one GSE, or None when GEO has no such public series."""
    s = http.get(f"{EUTILS}/esearch.fcgi", {"db": "gds", "term": f"{gse}[ACCN] AND gse[ETYP]",
                                               "retmode": "json"})
    try:
        ids = json.loads(s)["esearchresult"]["idlist"] if s else []
    except (ValueError, KeyError):
        ids = []
    if not ids:
        return None
    s = http.get(f"{EUTILS}/esummary.fcgi", {"db": "gds", "id": ids[0], "retmode": "json"})
    try:
        rec = json.loads(s)["result"][ids[0]]
    except (ValueError, KeyError, TypeError):
        return None
    text = " ".join(str(rec.get(k, "")) for k in ("title", "summary"))
    return {
        "gse": gse,
        "title": rec.get("title", ""),
        "summary": rec.get("summary", ""),
        "gdstype": rec.get("gdstype", ""),
        "data_type": classify_gdstype(rec.get("gdstype", ""), text),
        "taxon": rec.get("taxon", ""),
        "n_samples": int(rec.get("n_samples") or 0),
        "platforms": rec.get("gpl", ""),
        "suppfile": rec.get("suppfile", ""),
        "sample_titles": [x.get("title", "") for x in rec.get("samples", []) or []],
        "pubmedids": [str(x) for x in rec.get("pubmedids", []) or []],
    }


def geo_linked_series(pmid: str, http: Http) -> list[str]:
    """GSE accessions GEO links to a PubMed ID (elink pubmed -> gds)."""
    if not pmid or not http:
        return []
    s = http.get(f"{EUTILS}/elink.fcgi", {"dbfrom": "pubmed", "db": "gds", "id": pmid, "retmode": "json"})
    try:
        sets = json.loads(s)["linksets"] if s else []
    except (ValueError, KeyError):
        return []
    uids = [u for ls in sets for db in ls.get("linksetdbs", []) if db.get("linkname") == "pubmed_gds"
            for u in db.get("links", [])]
    # A GDS uid 2xxxxxxxx is series GSExxxxxxxx; other uids are DataSets and platforms.
    return sorted({f"GSE{int(u) - 200000000}" for u in uids if str(u).startswith("2") and len(str(u)) == 9},
                  key=lambda g: int(g[3:]))


def geo_owned_series(row: dict, http: Http) -> list[str]:
    """Series GEO links to this paper first: the paper's own deposit even when its text
    omits the accession (GSE235174), calls it "publicly available", or was not fetched.

    A series GEO also links to an earlier paper belongs to that paper (GSE163356, reused
    for validation); a preprint and its journal version count as one paper."""
    pmid = str(row.get("pmid") or "").strip()
    own = []
    for gse in geo_linked_series(pmid, http):
        info = geo_summary(gse, http)
        first = (info or {}).get("pubmedids", [])[:1]
        if info and (first == [pmid] or _norm_title(info["title"]) == _norm_title(row.get("title", ""))
                     or _same_paper(row, first, http)):
            own.append(gse)
    return own


# ---------------------------------------------------------------------------
# Layer 3: supplementary files (bulk RNA-seq)
# ---------------------------------------------------------------------------

def list_geo_dir(url: str, http: Http) -> list[str]:
    html = http.get(url)
    if not html:
        return []
    names = re.findall(r'href="([^"?/][^"]*)"', html)
    return [n for n in dict.fromkeys(names) if not n.startswith(("http", "/", "?"))]


def classify_supplementary(names: list[str]) -> dict:
    """What the supplementary files offer for a count-based analysis.

    status: ``raw_counts`` (a file named like a count matrix), ``raw_tar`` (only a
    per-sample ``_RAW.tar``; its members may be counts or FPKM and need inspection),
    ``normalized_only`` (FPKM/TPM/normalised files only), ``unlabelled_matrix`` (a table
    whose name says neither), ``de_tables_only`` (only differential-expression results),
    ``none``.
    """
    names = [n for n in names if n.lower() not in ("filelist.txt",)]
    de = [n for n in names if _MATRIX_EXT.search(n) and _DE_NAME.search(n)]
    raw = [n for n in names if _MATRIX_EXT.search(n) and _RAW_COUNT_NAME.search(n)
           and not _NORMALISED_NAME.search(n) and n not in de]
    norm = [n for n in names if _MATRIX_EXT.search(n) and _NORMALISED_NAME.search(n) and n not in de]
    tar = [n for n in names if n.lower().endswith("_raw.tar")]
    other = [n for n in names if _MATRIX_EXT.search(n) and n not in raw and n not in norm and n not in de]
    if raw:
        status = "raw_counts"
    elif tar:
        status = "raw_tar"
    elif norm:
        status = "normalized_only"
    elif other:
        status = "unlabelled_matrix"
    elif de:
        status = "de_tables_only"
    else:
        status = "none"
    return {"status": status, "raw": raw, "normalized": norm, "tar": tar, "other": other, "de": de}


# ---------------------------------------------------------------------------
# Layer 4: sample design from the series matrix header
# ---------------------------------------------------------------------------

def parse_series_matrix_samples(text: str):
    """Sample table (pandas DataFrame) from the ``!Sample_`` header lines of a series matrix."""
    import pandas as pd
    cols: dict[str, list] = {}
    char_n: Counter = Counter()
    for line in text.splitlines():
        if line.startswith("!series_matrix_table_begin"):
            break
        if not line.startswith("!Sample_"):
            continue
        parts = next(csv.reader([line], delimiter="\t"))
        key, vals = parts[0][len("!Sample_"):], [v.strip() for v in parts[1:]]
        if key == "characteristics_ch1":
            # "tissue: muscle" -> column characteristics_ch1.<i>.tissue
            label = next((v.split(":", 1)[0].strip() for v in vals if ":" in v), "")
            char_n[key] += 1
            key = f"characteristics_ch1.{char_n[key]}.{label}" if label else f"{key}.{char_n[key]}"
            vals = [v.split(":", 1)[1].strip() if ":" in v else v for v in vals]
        elif key in cols:
            continue
        cols[key] = vals
    if not cols:
        return pd.DataFrame()
    n = max(len(v) for v in cols.values())
    return pd.DataFrame({k: v + [""] * (n - len(v)) for k, v in cols.items()})


def design_from_samples(df) -> dict:
    """Groups and the smallest group size, using metadata_structural's design-column picker
    (falls back to sample titles that differ only by a replicate number)."""
    from tools import metadata_structural as ms
    if df is None or df.empty:
        return {"status": "no_metadata"}
    exclude = {c for c in df.columns if c in ("geo_accession", "title")}
    col = ms._pick_design_column(df, exclude)
    if col is not None:
        series = df[col].astype(str)
        source = col
    else:
        series = ms._title_design_series(df, "title" if "title" in df.columns else None)
        source = "title"
    sex_col = next((c for c in df.columns if re.search(r"\b(sex|gender)\b", str(c), re.I)), "")
    sexes = sorted(set(df[sex_col].astype(str).str.lower())) if sex_col else []
    sex = df[sex_col].astype(str).map(_sex_of) if sex_col else None
    composition = _sex_composition(sex)
    if series is None:
        return {"status": "single_group", "n_samples": len(df), "sex_column": sex_col, "sexes": sexes,
                "sex_composition": composition}
    series = series.astype(str).str.strip()
    labelled = series[series != ""]
    sizes = labelled.value_counts()
    if len(sizes) < 2:
        return {"status": "single_group", "n_samples": len(df), "sex_column": sex_col, "sexes": sexes,
                "sex_composition": composition}
    out = {
        "status": "ok" if _two_arms_of_three(sizes) else "group_too_small",
        "design_column": source,
        "groups": {str(k): int(v) for k, v in sizes.items()},
        "min_group": int(sizes.min()),
        "n_samples": len(df),
        "sex_column": sex_col,
        "sexes": sexes,
        "sex_composition": composition,
    }
    # A series holding several experiments (GSE297515: exercise, injury and a 2-sample
    # Yoda1 arm; 67 samples unlabelled in this column) keeps its usable comparison.
    small = {str(k): int(v) for k, v in sizes.items() if v < 3}
    if small:
        out["small_groups"] = small
    if len(series) - len(labelled):
        out["unlabelled_samples"] = int(len(series) - len(labelled))
    if composition == "mixed":
        # Male and female are analysed separately: each sex needs its own two arms of three.
        per = {sx: {str(k): int(v) for k, v in labelled[sex[labelled.index] == sx].value_counts().items()}
               for sx in ("male", "female")}
        out["per_sex_groups"] = per
        out["per_sex_ok"] = {sx: _two_arms_of_three(g) for sx, g in per.items()}
    return out


# Group labels of an exercise arm ("AEX.d1", "Exercise-trained (ExT)", "wheel running").
_EX_ARM = re.compile(r"exercis|train|run|wheel|swim|treadmill|endurance|resistance|HIIT|sprint|"
                     r"(?<![a-z])(a?ex|ext)(?![a-z])", re.I)


def _two_arms_of_three(sizes) -> bool:
    """True when two groups that would be compared both have >= 3 samples: an exercise arm
    and another arm when the labels name exercise, else any two groups (pre/post, time
    points). A small unrelated group does not fail the series."""
    sizes = {k: int(v) for k, v in dict(sizes).items()}
    ex = [g for g in sizes if _EX_ARM.search(str(g))]
    if ex and len(ex) < len(sizes):
        return any(sizes[g] >= 3 for g in ex) and any(v >= 3 for g, v in sizes.items() if g not in ex)
    return sum(v >= 3 for v in sizes.values()) >= 2


def _sex_of(v: str) -> str:
    low = str(v).strip().lower()
    if re.fullmatch(r"(f|female|females|woman|women|girl|girls)", low):
        return "female"
    if re.fullmatch(r"(m|male|males|man|men|boy|boys)", low):
        return "male"
    return ""


def _sex_composition(sex) -> str:
    """all_male / all_female / mixed, or unknown (no usable sex column: check X/Y genes)."""
    if sex is None:
        return "unknown"
    seen = set(sex) - {""}
    if seen == {"male"}:
        return "all_male"
    if seen == {"female"}:
        return "all_female"
    return "mixed" if seen == {"male", "female"} else "unknown"


def geo_design(gse: str, http: Http) -> dict:
    stub = geo_stub(gse)
    files = [n for n in list_geo_dir(GEO_MATRIX.format(stub=stub, gse=gse), http)
             if n.endswith("_series_matrix.txt.gz")]
    if not files:
        return {"status": "no_series_matrix"}
    raw = http.get(GEO_MATRIX.format(stub=stub, gse=gse) + files[0], binary=True, timeout=120)
    if not raw:
        return {"status": "no_series_matrix"}
    # Only the header is needed; stop decompressing at the data table.
    head = []
    with gzip.open(io.BytesIO(raw), "rt", encoding="utf-8", errors="replace") as fh:
        for line in fh:
            if line.startswith("!series_matrix_table_begin"):
                break
            head.append(line)
    samples = parse_series_matrix_samples("".join(head))
    out = design_from_samples(samples)
    out["matrix_files"] = len(files)
    # Every data-processing line (the sample table keeps only the first): the one that
    # names RPKM or salmon is often the last.
    proc = [next(csv.reader([ln], delimiter="\t"))[1:2] for ln in head if ln.startswith("!Sample_data_processing")]
    out["processing"] = " | ".join(dict.fromkeys(v[0].strip() for v in proc if v and v[0].strip()))[:2000]
    # Per-sample supplementary files (the members of a _RAW.tar), first file per sample.
    cols = [c for c in samples.columns if c.startswith("supplementary_file")]
    out["sample_files"] = [u for u in (samples[c].iloc[i] for i in range(len(samples)) for c in cols[:1])
                           if u and u.upper() != "NONE"]
    return out


# ---------------------------------------------------------------------------
# Layer 3b: what is inside a _RAW.tar, from the per-sample files
# ---------------------------------------------------------------------------

# Members of a _RAW.tar that hold one sample's table: salmon quant.sf, RSEM *.results too.
_MEMBER_EXT = re.compile(r"\.(csv|tsv|txt|tab|sf|results)(\.gz)?$", re.I)
# 10x / single-cell members: a series GEO labels bulk RNA-seq can still be single-cell.
_SINGLE_CELL_FILE = re.compile(r"barcodes\.tsv|features\.tsv|matrix\.mtx|feature_bc_matrix|\.h5ad?(\.gz)?$", re.I)
_NON_MATRIX_EXT = re.compile(
    r"\.(bw|bigwig|bed|bedgraph|bdg|wig|bam|sam|cram|h5|hdf5|mtx|narrowpeak|broadpeak|cel|idat|"
    r"gtf|gff3?|fa|fasta|fastq|fq|vcf|tdf|png|pdf|cov|rds|rda)(\.gz)?$", re.I)


def _numeric_columns(text: str, max_rows: int = 3000) -> tuple[list[str], list[list[float]]]:
    """(header, numeric columns) of a delimited per-sample file; '#' lines are skipped."""
    lines = [ln for ln in text.splitlines() if ln.strip() and not ln.startswith("#")][:max_rows + 1]
    if not lines:
        return [], []
    sep = "\t" if "\t" in lines[0] else ("," if "," in lines[0] else None)
    rows = [ln.split(sep) for ln in lines]
    header = []
    if not any(_is_number(v) for v in rows[0][1:]):
        header, rows = [h.strip().strip('"') for h in rows[0]], rows[1:]
    width = max((len(r) for r in rows), default=0)
    cols = []
    for j in range(1, width):
        vals = [r[j] for r in rows if len(r) > j]
        nums = [float(v.strip().strip('"')) for v in vals if _is_number(v)]
        cols.append(nums if vals and len(nums) >= 0.9 * len(vals) else [])
    return header, cols


def _is_number(v: str) -> bool:
    try:
        float(v.strip().strip('"'))
        return True
    except ValueError:
        return False


def sniff_count_file(text: str) -> str:
    """'raw_counts', 'normalized' or 'unknown' for one per-sample expression file."""
    header, cols = _numeric_columns(text)
    low = [h.lower() for h in header]
    # RSEM *.genes.results (expected_count), salmon quant.sf (NumReads), kallisto
    # abundance.tsv (est_counts): estimated counts sit next to TPM/FPKM.
    if any(h in low for h in ("expected_count", "numreads", "est_counts")):
        return "raw_counts"
    numeric = [(j, c) for j, c in enumerate(cols) if c]
    if not numeric:
        return "unknown"
    j, col = numeric[-1]   # featureCounts/HTSeq/STAR: the count is the rightmost number
    name = header[j + 1].lower() if len(header) > j + 1 else ""
    if _NORMALISED_NAME.search(name):
        return "normalized"
    if len(col) < 50:
        return "unknown"
    integer = sum(1 for v in col if v >= 0 and v == int(v)) / len(col)
    return "raw_counts" if integer >= 0.99 else "normalized"


def sniff_matrix_file(text: str) -> str:
    """'raw_counts', 'mostly_integer' or 'decimal' (or 'unknown') for a genes x samples
    matrix, from its values.

    File names are not enough: ``GSE273041_raw_counts.txt.gz`` holds RPKM. Decimals alone
    do not prove normalisation either: salmon/RSEM/Partek estimated counts and
    featureCounts fractional multi-mapper counts are decimal counts. ``check_raw_values``
    settles a decimal file with the series' data-processing text."""
    header, cols = _numeric_columns(text)
    if any(h.lower() in ("expected_count", "numreads", "est_counts") or _COUNT_COLUMN.search(h) for h in header):
        return "raw_counts" if _count_columns_integer(header, cols) else "decimal"
    values = [v for c in cols for v in c]
    if len(values) < 200:
        return "unknown"
    # Within 0.01 of an integer: salmon's merged counts carry 132.001-style noise.
    integer = sum(1 for v in values if v >= 0 and abs(v - round(v)) <= 0.01) / len(values)
    if integer >= 0.99:
        return "raw_counts"
    # A file holding blocks of several kinds (mRNA counts, lncRNA FPKM, circRNA counts):
    # two or more integer sample columns are a count block.
    names = [header[j + 1] if len(header) > j + 1 else "" for j in range(len(cols))]
    per = [(n, _integer_share(c)) for n, c in zip(names, cols) if len(c) >= 20]
    whole = [f for n, f in per if f >= 0.99 and not _NOT_SAMPLE_COLUMN.search(n)]
    if len(whole) >= 2 and any(f < 0.5 for _, f in per):
        return "mixed"
    return "mostly_integer" if integer >= 0.5 else "decimal"


# A per-sample count column in a wide table ("Count (FC1)", "read_count_S1"), not a
# normalised one ("normalized count").
_COUNT_COLUMN = re.compile(r"^(read[ _]?)?counts?\b|\b(read[ _]?)?counts?[ _]*\(", re.I)
_NOT_SAMPLE_COLUMN = re.compile(r"start|end|length|len\b|pos|chr|strand|width|id$", re.I)


def _integer_share(col: list) -> float:
    return sum(1 for v in col if v >= 0 and abs(v - round(v)) <= 0.01) / len(col) if col else 0.0


def _count_columns_integer(header: list, cols: list) -> bool:
    """True when the columns named like counts (Novogene "Count (S1)", RSEM expected_count)
    hold integer-like values."""
    picked = [c for j, c in enumerate(cols) if c and len(header) > j + 1 and
              (_COUNT_COLUMN.search(header[j + 1]) or header[j + 1].lower() in ("expected_count", "numreads",
                                                                                "est_counts"))
              and not _NORMALISED_NAME.search(header[j + 1].replace("count", ""))]
    return bool(picked) and sum(_integer_share(c) >= 0.99 for c in picked) >= len(picked) / 2


def _head_text(raw: bytes) -> str:
    """Text of a (possibly truncated, possibly gzipped) file head, last partial line dropped."""
    import zlib
    if raw[:2] == b"\x1f\x8b":
        try:
            raw = zlib.decompressobj(16 + zlib.MAX_WBITS).decompress(raw)
        except zlib.error:
            return ""
    text = raw.decode("utf-8", errors="replace")
    return text.rsplit("\n", 1)[0] if "\n" in text else text


# A unit, not "normalized counts": GEO texts often list a raw and a normalised file.
_PROCESSING_NORMALISED = re.compile(r"fpkm|rpkm|tpm|(?<![a-z])cpm|per kilobase", re.I)
_PROCESSING_ESTIMATED = re.compile(r"rsem|salmon|kallisto|expected[_ ]count|estimated counts?|tximport|partek e/?m|"
                                   r"fraction(al)?|multi-?mapp", re.I)


def _sniff_supplementary(gse: str, name: str, http: Http, max_bytes: int = 2_000_000) -> str:
    """sniff_matrix_file on one supplementary file: the head of a text file, or the first
    rows of an Excel workbook (read whole; Novogene's gene_expression.xlsx is ~15 MB)."""
    url = GEO_SUPPL.format(stub=geo_stub(gse), gse=gse) + name
    if not re.search(r"\.xlsx?$", name, re.I):
        raw = http.get(url, binary=True, timeout=120, max_bytes=max_bytes)
        return sniff_matrix_file(_head_text(raw)) if raw else "unknown"
    if not name.lower().endswith(".xlsx"):
        return "unknown"
    raw = http.get(url, binary=True, timeout=300)
    if not raw or len(raw) > 60_000_000:
        return "unknown"
    try:
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(max_row=3001, values_only=True)]
    except Exception:
        return "unknown"
    return sniff_matrix_file("\n".join("\t".join(r) for r in rows if any(r)))


def check_raw_values(gse: str, files: dict, http: Http, processing: str = "",
                     max_bytes: int = 2_000_000, key: str = "raw") -> dict:
    """Sniff the values of the files named like raw counts (first two text files).

    ``key="other"`` checks the unlabelled tables instead. Returns
    ``{"status": ..., "sniffed": {...}}`` with status ``raw_counts`` (integers, or an
    integer count block in a mixed file),
    ``estimated_counts`` (decimal counts: salmon, RSEM, fractional featureCounts),
    ``normalized_only`` (decimals, and the data-processing text names FPKM/RPKM/TPM),
    ``decimal_unexplained`` (decimals the processing text does not explain) or ``unknown``
    (no text file could be read: Excel, download failure)."""
    sniffed = {name: _sniff_supplementary(gse, name, http, max_bytes) for name in files.get(key, [])[:2]}
    kinds = set(sniffed.values()) - {"unknown"}
    if not kinds:
        status = "unknown"
    elif "raw_counts" in kinds or "mixed" in kinds:
        status = "raw_counts"
    elif kinds == {"mostly_integer"}:
        # Mostly whole numbers are counts (estimated or fractional), whatever the
        # processing text says about the series' other, normalised files.
        status = "estimated_counts"
    elif _PROCESSING_NORMALISED.search(processing) and not _PROCESSING_ESTIMATED.search(processing):
        status = "normalized_only"
    elif _PROCESSING_ESTIMATED.search(processing) or kinds == {"mostly_integer"}:
        status = "estimated_counts"
    else:
        status = "decimal_unexplained"
    return {"status": status, "sniffed": sniffed}


def inspect_sample_files(urls: list[str], http: Http, n_sniff: int = 2) -> dict:
    """What a series' per-sample files (the members of its _RAW.tar) hold.

    status: ``raw_counts`` (per-sample count tables), ``normalized_only`` (FPKM/TPM or
    non-integer values), ``no_count_files`` (only tracks, peaks, BAMs, 10x matrices...),
    or ``unknown`` (could not tell; stays needs_inspection).
    """
    names = [u.rsplit("/", 1)[-1] for u in urls]
    if not names:
        return {"status": "unknown", "why": "no per-sample files listed"}
    if any(_SINGLE_CELL_FILE.search(n) for n in names):
        return {"status": "single_cell", "examples": names[:3]}
    tables = [(u, n) for u, n in zip(urls, names) if _MEMBER_EXT.search(n) and not _NON_MATRIX_EXT.search(n)]
    if not tables:
        return {"status": "no_count_files", "examples": names[:3]}
    # kallisto's abundance.tsv carries est_counts: let the values decide.
    if all(_NORMALISED_NAME.search(n) and not _RAW_COUNT_NAME.search(n) and "abundance.tsv" not in n.lower()
           for _, n in tables):
        return {"status": "normalized_only", "examples": [n for _, n in tables[:3]], "why": "file names"}
    kinds = []
    for u, n in tables[:n_sniff]:
        raw = http.get(u.replace("ftp://", "https://", 1), binary=True, timeout=120)
        if not raw:
            continue
        try:
            text = (gzip.decompress(raw) if raw[:2] == b"\x1f\x8b" else raw).decode("utf-8", errors="replace")
        except (OSError, EOFError):
            continue
        kinds.append(sniff_count_file(text))
    examples = [n for _, n in tables[:3]]
    if kinds and all(k == "raw_counts" for k in kinds):
        return {"status": "raw_counts", "examples": examples, "n_tables": len(tables)}
    if kinds and all(k == "normalized" for k in kinds):
        return {"status": "normalized_only", "examples": examples, "why": "values"}
    return {"status": "unknown", "examples": examples, "sniffed": kinds}


# ---------------------------------------------------------------------------
# Verdict
# ---------------------------------------------------------------------------

SUPPORTED_ROUTES = {"rnaseq"}   # raw-count-only run; widen to route other data types


@dataclass
class Screen:
    paper_id: str
    title: str = ""
    pmcid: str = ""
    doi: str = ""
    text_source: str = ""
    own_gse: list = field(default_factory=list)
    other_repos: dict = field(default_factory=dict)
    cited_gse: list = field(default_factory=list)
    evidence: dict = field(default_factory=dict)
    series: list = field(default_factory=list)
    data_types: str = ""
    verdict: str = ""
    reason: str = ""
    analyzable_gse: list = field(default_factory=list)
    source: str = ""
    query: str = ""
    linked_gse: str = ""
    route: str = ""
    missing: str = ""
    flags: list = field(default_factory=list)

    def row(self) -> dict:
        return {
            "paper_id": self.paper_id, "title": self.title, "pmcid": self.pmcid, "doi": self.doi,
            "source": self.source, "query": self.query, "linked_gse": self.linked_gse,
            "verdict": self.verdict, "reason": self.reason, "route": self.route,
            "missing": self.missing, "flags": ";".join(self.flags),
            "analyzable_gse": ";".join(self.analyzable_gse), "data_types": self.data_types,
            "own_gse": ";".join(self.own_gse), "cited_gse": ";".join(self.cited_gse),
            "other_repos": ";".join(f"{k}:{','.join(v)}" for k, v in self.other_repos.items()),
            "text_source": self.text_source,
            "series_json": json.dumps(self.series, ensure_ascii=False),
            "evidence_json": json.dumps(self.evidence, ensure_ascii=False),
        }


# Analysis route per matrix (references/matrix-method-policy.md of the workflow skill).
ROUTE = {"raw_counts": "counts", "raw_counts_in_tar": "counts", "normalized_only": "limma"}

# What each data type needs before it can be analysed. Only bulk RNA-seq is routed now;
# the other rows say what a later route would need.
REQUIRED_DATA = {
    "bulk RNA-seq, counts (DESeq2/edgeR/limma-voom)": "genes x samples raw count matrix, or one count table "
        "per sample (featureCounts/HTSeq/STAR, RSEM expected_count, salmon NumReads); sample groups; sex per "
        "sample or X/Y genes in the matrix",
    "bulk RNA-seq, TPM/FPKM/log (limma)": "genes x samples TPM/FPKM/CPM or log matrix; sample groups; sex per "
        "sample or X/Y genes in the matrix",
    "microarray (not routed)": "normalised log expression matrix (series matrix or CEL files) with the platform "
        "annotation; sample groups; sex",
    "single-cell (not routed)": "raw UMI counts per cell with a cell-to-sample map, for pseudobulk; sample groups; sex",
    "methylation (not routed)": "beta or M values per CpG; sample groups; sex",
}

# What a paper lacks, by the layer it failed.
MISSING = {
    "no_full_text": "full text, or a GEO series linked to the paper's PMID",
    "off_topic": "an exercise topic",
    "no_exercise_contrast": "the paper's own exercise groups or protocol",
    "no_own_deposit": "the paper's own GEO deposit",
    "data_on_request": "a public deposit (data on request only)",
    "non_geo_repository": "a GEO series with a processed matrix (data elsewhere or raw reads only)",
    "geo_not_public": "a public GEO series",
    "unsupported_data_type": "bulk RNA-seq data",
    "unsupported_organism": "human, mouse or rat samples",
    "no_processed_matrix": "a count or TPM/FPKM matrix (GEO has no processed matrix)",
    "de_tables_only": "a count or TPM/FPKM matrix (GEO has DE tables only)",
    "normalized_only": "raw counts (only TPM/FPKM; limma is off with --counts-only)",
    "single_group": "two or more sample groups",
    "group_too_small": "3 or more samples in each of two compared groups",
}


def series_flags(design: dict | None) -> list[str]:
    """Checks for the analysis: sex composition (all male / all female / mixed / unknown),
    arms short of 3 once split by sex, extra small groups, unlabelled samples."""
    if not design:
        return []
    flags = [f"sex:{design.get('sex_composition', 'unknown')}"]
    for sx, ok in (design.get("per_sex_ok") or {}).items():
        if not ok:
            flags.append(f"{sx}_arms_below_3")
    if design.get("small_groups"):
        flags.append("small_groups:" + ",".join(f"{k}={v}" for k, v in design["small_groups"].items()))
    if design.get("unlabelled_samples"):
        flags.append(f"unlabelled_samples:{design['unlabelled_samples']}")
    return flags


# Order matters: the reason recorded for a paper is the earliest layer it fails.
REASONS = (
    "no_full_text", "off_topic", "no_exercise_contrast", "no_own_deposit", "data_on_request", "non_geo_repository",
    "geo_not_public", "unsupported_data_type", "unsupported_organism", "no_processed_matrix",
    "de_tables_only", "normalized_only",
    "single_group", "group_too_small",
)


def organism_ok(taxon: str, organisms: set | None) -> bool:
    """True when one of the series' organisms is supported (``organisms=None``: any)."""
    if not organisms:
        return True
    # By genus: GEO labels some laboratory rats "Rattus rattus" instead of R. norvegicus.
    genera = {o.split()[0] for o in organisms}
    return any(t.strip().lower().split(" ")[0] in genera for t in re.split(r"[;,]", taxon or "") if t.strip())


def judge_series(info: dict | None, files: dict | None, design: dict | None,
                 routes: set = SUPPORTED_ROUTES, organisms: set | None = None,
                 counts_only: bool = False) -> tuple[str, str]:
    """(verdict, reason) for one GSE. TPM/FPKM/log matrices pass (they go to limma)
    unless ``counts_only``."""
    if info is None:
        return "excluded", "geo_not_public"
    if not set(info["data_type"].split("+")) & routes:
        return "excluded", f"unsupported_data_type:{info['data_type']}"
    if not organism_ok(info.get("taxon", ""), organisms):
        return "excluded", f"unsupported_organism:{info.get('taxon', '')}"
    if files is not None:
        st = files["status"]
        if st == "single_cell":
            return "excluded", "unsupported_data_type:scrnaseq"
        if st == "none":
            return "excluded", "no_processed_matrix"
        if st == "de_tables_only":
            return "excluded", "de_tables_only"
        if st == "normalized_only" and counts_only:
            return "excluded", "normalized_only"
        if st == "decimal_unexplained":
            return "needs_inspection", "matrix_decimal_unexplained"
    if design is not None:
        ds = design.get("status")
        if ds == "single_group":
            return "excluded", "single_group"
        if ds == "group_too_small":
            return "excluded", "group_too_small"
    if files is not None and files["status"] in ("raw_tar", "unlabelled_matrix"):
        return "needs_inspection", f"matrix_{files['status']}"
    if design is not None and design.get("status") not in ("ok",):
        return "needs_inspection", f"design_{design.get('status')}"
    return "analyzable", ""


def on_topic(row: dict, doc: dict, topic: re.Pattern | None) -> bool:
    """True when the paper's title or abstract mentions the topic (``topic=None``: any).

    Literature and GEO searches match on any field, so they return papers that only
    mention the topic in passing (a Trichomonas drug-resistance study, an E. coli biofilm).
    """
    if topic is None:
        return True
    blocks = doc.get("blocks", [])
    abstract = [b["text"] for b in blocks if b.get("section_class") == "abstract"] or [b["text"] for b in blocks[:5]]
    return bool(topic.search(" ".join([row.get("title", ""), doc.get("title", "")] + abstract)))


# Sections where a paper describes its own design; the introduction and discussion
# talk about exercise in general.
_DESIGN_SECTIONS = {"abstract", "methods", "results", "other"}


def contrast_evidence(doc: dict, contrast: re.Pattern | None, limit: int = 4) -> list[str]:
    """Phrases where the paper describes its own exercise groups, protocol or comparison.

    A topic word is not enough: an aging study that measured treadmill capacity
    (GSE152133), a diet study and a vessel-stretch study all name exercise in the
    abstract or introduction but have no exercise arm. ``contrast=None``: not checked."""
    if contrast is None:
        return ["(not checked)"]
    hits = []
    for b in doc.get("blocks", []):
        if b.get("section_class") in _DESIGN_SECTIONS:
            hits += [f"{b.get('id', '')}: {m.group(0)[:80]}" for m in contrast.finditer(b.get("text", ""))]
            if len(hits) >= limit:
                break
    return hits[:limit]


def _norm_title(t: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", (t or "").lower())


def _same_paper(row: dict, pmids: list[str], http: Http) -> bool:
    """True when one of ``pmids`` is this paper under another record: a bioRxiv preprint
    and its journal version get different PMIDs but keep the title (BMAL1 and exercise
    training: PMID 37905004 vs 38950777)."""
    title = _norm_title(row.get("title", ""))
    if not title or not http:
        return False
    return any(_norm_title(r["title"]) == title for r in _pubmed_summaries(pmids, http, "geo_link"))


def screen_paper(row: dict, doc: dict | None, http: Http | None, check_files: bool = True,
                 check_design: bool = True, routes: set = SUPPORTED_ROUTES,
                 organisms: set | None = None, topic: re.Pattern | None = None,
                 contrast: re.Pattern | None = None, counts_only: bool = False) -> Screen:
    """Screen one paper. ``doc`` is its structured full text (None when unavailable).

    ``organisms`` (lower-case species names), ``topic`` (a regex for the title or
    abstract) and ``contrast`` (a regex for the paper's own exercise groups or protocol,
    see ``contrast_evidence``) are off when None; the CLI turns them on by default.
    Series GEO links to the paper first count as its own, with or without full text.
    ``counts_only`` excludes TPM/FPKM-only series instead of routing them to limma."""
    s = _screen_paper(row, doc, http, check_files, check_design, routes, organisms, topic, contrast, counts_only)
    best = next((r for r in s.series if r["gse"] in s.analyzable_gse), None)
    if best is None:
        best = next((r for r in s.series if r["verdict"] == "needs_inspection"), None)
    if best is not None:
        s.flags = series_flags(best.get("design"))
    s.route = ";".join(dict.fromkeys(r.get("route", "") for r in s.series if r["gse"] in s.analyzable_gse))
    if s.verdict == "excluded":
        s.missing = MISSING.get(s.reason.split(":")[0], "")
    elif s.verdict == "analyzable" and "sex:unknown" in s.flags:
        s.missing = "sex per sample (check X/Y genes)"
    return s


def _screen_paper(row, doc, http, check_files, check_design, routes, organisms, topic, contrast,
                  counts_only) -> Screen:
    s = Screen(paper_id=row.get("paper_id") or row.get("paperId") or row.get("pmcid") or row.get("doi", ""),
               title=row.get("title", ""), pmcid=row.get("pmcid", ""), doi=row.get("doi", ""),
               source=row.get("source", ""), query=row.get("query", ""), linked_gse=row.get("linked_gse", ""))
    geo_own = geo_owned_series(row, http) if http else []
    if not doc or not doc.get("blocks"):
        if not geo_own:
            s.verdict, s.reason = "excluded", "no_full_text"
            return s
        # No full text, but GEO links the paper to its own series: judge the topic and
        # the exercise design from the title and the series' GEO summaries.
        infos = [geo_summary(g, http) or {} for g in geo_own]
        geo_text = " ".join(f"{i.get('title', '')}. {i.get('summary', '')}" for i in infos)
        doc = {"title": row.get("title", ""), "structure_quality": "geo_only",
               "blocks": [{"id": "GEO", "section_class": "abstract", "kind": "paragraph", "text": geo_text}]}
    s.text_source = doc.get("structure_quality", "")
    if not on_topic(row, doc, topic):
        s.verdict, s.reason = "excluded", "off_topic"
        return s
    found = contrast_evidence(doc, contrast)
    if not found:
        s.verdict, s.reason = "excluded", "no_exercise_contrast"
        return s
    dep = deposits_from_doc(doc)
    for gse in geo_own:
        # GEO's own link wins over the wording around a mention (or its absence).
        if gse not in dep["own"].get("geo", []):
            dep["own"].setdefault("geo", []).append(gse)
        if gse in dep["cited"]:
            dep["cited"].remove(gse)
        dep["said_own"].append(gse)
    s.own_gse = dep["own"].get("geo", [])
    s.other_repos = {k: v for k, v in dep["own"].items() if k != "geo"}
    s.cited_gse = dep["cited"]
    s.evidence = {k: v[:4] for k, v in dep["evidence"].items()}
    if not s.own_gse:
        if s.other_repos:
            s.verdict, s.reason = "excluded", "non_geo_repository:" + ",".join(sorted(s.other_repos))
        elif any(a.startswith("GSE") for a in s.cited_gse):
            # Analyses of other people's GEO series (meta-analyses, re-analyses); an
            # "upon request" line there is about derived results, not a hidden deposit.
            s.verdict, s.reason = "excluded", "no_own_deposit:reuses_public_gse"
        elif dep["on_request"]:
            s.verdict, s.reason = "excluded", "data_on_request"
        else:
            s.verdict, s.reason = "excluded", "no_own_deposit"
        return s
    results = []
    types = []
    pmid = str(row.get("pmid") or "").strip()
    for gse in s.own_gse:
        info = geo_summary(gse, http) if http else None
        files = design = None
        if (info and pmid and info.get("pubmedids") and pmid not in info["pubmedids"]
                and gse not in dep["said_own"] and not _same_paper(row, info["pubmedids"], http)):
            # GEO links the series to other papers only, and this one never says it
            # deposited it: a public dataset listed in Methods (GSE91061 in a radiomics
            # paper). A paper GEO has not linked yet has an empty list and passes.
            results.append({"gse": gse, "verdict": "excluded", "reason": "no_own_deposit:linked_to_other_paper",
                            "data_type": info["data_type"], "taxon": info["taxon"],
                            "n_samples": info["n_samples"], "files": None, "design": None,
                            "linked_pmids": info["pubmedids"]})
            types.append(info["data_type"])
            continue
        if info and set(info["data_type"].split("+")) & routes and organism_ok(info["taxon"], organisms):
            if check_files:
                files = classify_supplementary(list_geo_dir(GEO_SUPPL.format(stub=geo_stub(gse), gse=gse), http))
            if check_design:
                design = geo_design(gse, http)
                if files and files["status"] == "raw_counts":
                    files["value_check"] = vc = check_raw_values(gse, files, http, design.get("processing", ""))
                    if vc["status"] in ("normalized_only", "decimal_unexplained"):
                        # Named raw_counts but holds decimals (GSE273041: RPKM). A _RAW.tar may
                        # still hold per-sample counts (GSE179302).
                        files["status"] = "raw_tar" if files["tar"] else vc["status"]
                    vc["processing"] = design["processing"][:300]
                elif files and files["status"] == "unlabelled_matrix":
                    # A table named neither counts nor FPKM: its columns or values decide
                    # (Novogene gene_expression.xlsx has "Count (S1)" columns).
                    files["value_check"] = vc = check_raw_values(gse, files, http, design.get("processing", ""),
                                                                 key="other")
                    files["status"] = {"raw_counts": "raw_counts", "estimated_counts": "raw_counts",
                                       "normalized_only": "normalized_only"}.get(vc["status"], "unlabelled_matrix")
                    vc["processing"] = design["processing"][:300]
                design.pop("processing", None)
                if files and files["status"] == "raw_tar":
                    files["tar_inspection"] = inside = inspect_sample_files(design.get("sample_files", []), http)
                    files["status"] = {"raw_counts": "raw_counts_in_tar", "normalized_only": "normalized_only",
                                       "no_count_files": "none", "single_cell": "single_cell"}.get(inside["status"],
                                                                                                  "raw_tar")
                design.pop("sample_files", None)
        verdict, reason = judge_series(info, files, design, routes, organisms, counts_only)
        types.append(info["data_type"] if info else "not_public")
        results.append({"gse": gse, "verdict": verdict, "reason": reason,
                        "route": ROUTE.get((files or {}).get("status"), "") if verdict == "analyzable" else "",
                        "data_type": info["data_type"] if info else "", "taxon": info["taxon"] if info else "",
                        "n_samples": info["n_samples"] if info else 0,
                        "files": files, "design": design})
    s.series = results
    s.data_types = ";".join(dict.fromkeys(types))
    s.analyzable_gse = [r["gse"] for r in results if r["verdict"] == "analyzable"]
    if s.analyzable_gse:
        s.verdict, s.reason = "analyzable", ""
    elif any(r["verdict"] == "needs_inspection" for r in results):
        s.verdict = "needs_inspection"
        s.reason = next(r["reason"] for r in results if r["verdict"] == "needs_inspection")
    else:
        s.verdict = "excluded"
        # Report the reason that got furthest through the layers.
        rank = {r: i for i, r in enumerate(REASONS)}
        s.reason = max((r["reason"] for r in results), key=lambda x: rank.get(x.split(":")[0], -1))
    return s


# ---------------------------------------------------------------------------
# Paper collection (Semantic Scholar + PubMed) and full-text loading
# ---------------------------------------------------------------------------

SOURCES = ("s2", "pubmed", "geo")
_GEO_RNASEQ = '"expression profiling by high throughput sequencing"[DataSet Type]'


def _s2_rows(query: str, http: Http, n: int, end_year: int = 0, start_year: int = 0) -> list[dict]:
    rows: list[dict] = []
    offset = 0
    while len(rows) < n and offset < 1000:
        params = {"query": query, "fields": "title,abstract,year,externalIds,openAccessPdf",
                  "limit": min(100, n - len(rows)), "offset": offset,
                  "fieldsOfStudy": "Medicine,Biology"}
        if end_year or start_year:
            params["year"] = f"{start_year or ''}-{end_year or ''}"
        s = http.get(S2_SEARCH + "?openAccessPdf", params)
        data = (json.loads(s).get("data") or []) if s else []
        if s is None:
            print("Semantic Scholar search failed (rate limited?); set S2_API_KEY for a steadier quota. "
                  "Continuing without it.", file=sys.stderr)
        if not data:
            break
        for p in data:
            ext = p.get("externalIds") or {}
            rows.append({"paper_id": p.get("paperId", ""), "source": "semantic_scholar",
                         "title": (p.get("title") or "").strip(), "year": p.get("year") or "",
                         "doi": (ext.get("DOI") or "").lower(), "pmid": str(ext.get("PubMed") or ""),
                         "pmcid": _pmcid(ext.get("PubMedCentral", "")),
                         "pdf_url": (p.get("openAccessPdf") or {}).get("url", "")})
        offset += len(data)
    return rows


def _pubmed_summaries(pmids: list[str], http: Http, source: str) -> list[dict]:
    rows = []
    for i in range(0, len(pmids), 200):
        s = http.get(f"{EUTILS}/esummary.fcgi", {"db": "pubmed", "id": ",".join(pmids[i:i + 200]),
                                                    "retmode": "json"})
        res = json.loads(s).get("result", {}) if s else {}
        for pid in res.get("uids", []):
            r = res[pid]
            ids = {a.get("idtype"): a.get("value", "") for a in r.get("articleids", [])}
            rows.append({"paper_id": f"PMID{pid}", "source": source, "title": r.get("title", "").strip(),
                         "year": (r.get("pubdate") or "")[:4], "doi": ids.get("doi", "").lower(),
                         "pmid": pid, "pmcid": _pmcid(ids.get("pmc", "")), "pdf_url": ""})
    return rows


def _pubmed_rows(query: str, http: Http, n: int, end_year: int = 0, start_year: int = 0) -> list[dict]:
    term = f"({query}) AND free full text[sb]"
    if end_year or start_year:
        term += f" AND {start_year or 1800}:{end_year or 3000}[dp]"
    s = http.get(f"{EUTILS}/esearch.fcgi", {"db": "pubmed", "term": term, "retmax": n,
                                               "sort": "relevance", "retmode": "json"})
    pmids = json.loads(s)["esearchresult"]["idlist"] if s else []
    return _pubmed_summaries(pmids, http, "pubmed")


def _geo_rows(query: str, http: Http, n: int, end_year: int = 0, start_year: int = 0) -> list[dict]:
    """Papers linked from GEO bulk RNA-seq series that match the query.

    Starting from the deposit finds papers that have their own data, which a literature
    search mostly does not (88 of 200 papers in the first run deposited nothing). GEO's
    PubMed link is usually the submitter's own paper; the ownership check still runs on
    the text. ``linked_gse`` records the series that led to the paper.
    """
    term = f"({query}) AND gse[Entry Type] AND {_GEO_RNASEQ}"
    if end_year or start_year:
        term += f" AND {start_year or 1800}:{end_year or 3000}[PDAT]"
    s = http.get(f"{EUTILS}/esearch.fcgi", {"db": "gds", "term": term, "retmax": max(n * 3, 100),
                                               "retmode": "json"})
    uids = json.loads(s)["esearchresult"]["idlist"] if s else []
    linked: dict[str, list] = {}
    for i in range(0, len(uids), 200):
        s = http.get(f"{EUTILS}/esummary.fcgi", {"db": "gds", "id": ",".join(uids[i:i + 200]),
                                                    "retmode": "json"})
        res = json.loads(s).get("result", {}) if s else {}
        for uid in res.get("uids", []):
            r = res[uid]
            for pid in r.get("pubmedids", []) or []:
                linked.setdefault(str(pid), []).append(r.get("accession", ""))
        if len(linked) >= n:
            break
    pmids = list(linked)[:n]
    rows = _pubmed_summaries(pmids, http, "geo")
    for r in rows:
        r["linked_gse"] = ";".join(dict.fromkeys(linked.get(r["pmid"], [])))
    return rows


def _paper_keys(r: dict) -> list[str]:
    return [k for k in (r.get("doi") and "doi:" + r["doi"], r.get("pmid") and "pmid:" + r["pmid"],
                        "t:" + re.sub(r"\W+", "", (r.get("title") or "").lower())[:80]) if k]


def collect_papers(query, http: Http, max_papers: int = 200, end_year: int = 0,
                   sources=("s2", "pubmed"), start_year: int = 0, shares: dict | None = None,
                   exclude: list[dict] | None = None) -> list[dict]:
    """Up to ``max_papers`` papers for one query or a list of queries.

    Each (query, source) list is relevance-ranked; the lists are taken in turn so every
    query and source is represented. Sources: ``s2`` (Semantic Scholar open-access hits),
    ``pubmed`` (PubMed free-full-text hits), ``geo`` (papers linked from matching GEO
    bulk RNA-seq series). Deduplicated by DOI, then PMID, then title; a duplicate fills in
    the kept row's missing PMCID / PDF URL / linked series.

    ``start_year``/``end_year`` bound the publication year. ``shares`` ({source: fraction})
    caps each source at that share of ``max_papers`` (a source that falls short leaves its
    room to the others). ``exclude`` rows (an earlier papers.csv) are never returned.
    """
    queries = [query] if isinstance(query, str) else list(query)
    fetch = {"s2": _s2_rows, "pubmed": _pubmed_rows, "geo": _geo_rows}
    lists = []
    for q in queries:
        for src in sources:
            want = max_papers if not shares else max(1, round(max_papers * shares.get(src, 0) * 3 / len(queries)))
            rows = fetch[src](q, http, want, end_year, start_year)
            for r in rows:
                r["query"] = q
                r.setdefault("linked_gse", "")
            lists.append(rows)
    merged = [x[i] for i in range(max((len(x) for x in lists), default=0)) for x in lists if i < len(x)]
    seen = {k for r in (exclude or []) for k in _paper_keys(r)}
    cap = ({s: round(max_papers * f) for s, f in shares.items()} if shares else None)
    kept: dict[str, dict] = {}
    out, overflow, used = [], [], Counter()
    for r in merged:
        keys = _paper_keys(r)
        if any(k in seen for k in keys):
            continue
        first = next((kept[k] for k in keys if k in kept), None)
        if first is not None:
            for f in ("pmcid", "pdf_url", "pmid", "doi", "linked_gse"):
                first[f] = first[f] or r[f]
            continue
        for k in keys:
            kept[k] = r
        src = {"semantic_scholar": "s2"}.get(r["source"], r["source"])
        if cap is not None and used[src] >= cap.get(src, 0):
            overflow.append(r)
            continue
        used[src] += 1
        out.append(r)
    out += overflow[:max(0, max_papers - len(out))]
    return out[:max_papers]


def _pmcid(v) -> str:
    v = str(v or "").strip().upper()
    return v if v.startswith("PMC") else (f"PMC{v}" if v.isdigit() else "")


def has_body(doc: dict, min_paragraphs: int = 5) -> bool:
    """True when a structured paper has text beyond its abstract and front matter."""
    n = sum(1 for b in doc.get("blocks", [])
            if b.get("kind") == "paragraph" and b.get("section_class") != "abstract")
    return n >= min_paragraphs


def load_full_text(row: dict, papers_dir: str, http: Http) -> dict | None:
    """Structured full text: cached .paper.json, else Europe PMC JATS, else the PDF."""
    from tools import paper_structure as ps
    stem = re.sub(r"[^A-Za-z0-9._-]", "_", row.get("pmcid") or row.get("paper_id") or "paper")
    cached = os.path.join(papers_dir, f"{stem}.paper.json")
    if os.path.exists(cached):
        return ps.load(cached)
    if row.get("pmcid"):
        # Europe PMC first; NCBI PMC (efetch) serves some OA and author-manuscript texts
        # Europe PMC answers with a 500. Both can return front matter only, so a text
        # counts only when it has a body.
        for url, params, label in (
                (ps.EPMC_FULLTEXT.format(pmcid=row["pmcid"]), None, "Europe PMC full-text XML"),
                (f"{EUTILS}/efetch.fcgi", {"db": "pmc", "id": row["pmcid"][3:], "retmode": "xml"},
                 "NCBI PMC efetch XML")):
            xml = http.get(url, params, binary=True)
            if not xml or b"<article" not in xml[:5000]:
                continue
            try:
                doc = ps.parse_jats(xml, source=f"{label} ({row['pmcid']})")
            except Exception:
                continue
            if has_body(doc):
                ps.save(doc, papers_dir, stem, raw_xml=xml)
                return doc
    if row.get("pdf_url"):
        pdf = http.get(row["pdf_url"], binary=True, timeout=120)
        if pdf and pdf[:5] == b"%PDF-":
            from tools import pdf_markdown
            os.makedirs(papers_dir, exist_ok=True)
            pdf_path = os.path.join(papers_dir, f"{stem}.pdf")
            with open(pdf_path, "wb") as f:
                f.write(pdf)
            doc, _ = pdf_markdown.save_pdf(pdf_path, papers_dir, stem, fetched_url=row["pdf_url"])
            return doc
    return None


# ---------------------------------------------------------------------------
# Batch + summary
# ---------------------------------------------------------------------------

FIELDS = ["paper_id", "title", "pmcid", "doi", "source", "query", "linked_gse", "verdict", "reason", "route",
          "missing", "flags", "analyzable_gse", "data_types",
          "own_gse", "cited_gse", "other_repos", "text_source", "series_json", "evidence_json"]


DEFAULT_ORGANISMS = {"homo sapiens", "mus musculus", "rattus norvegicus"}
# The exercise topic; --topic replaces it for another theme, --topic "" turns it off.
DEFAULT_TOPIC = (r"exercis|physical(ly)? (activ|train|fit|inactiv)|endurance|"
                 r"(exercise|endurance|resistance|strength|interval|aerobic|treadmill|swim(ming)?|physical|sprint|"
                 r"running|athletic|weight|muscle|sports?)[- ]train|\b(un)?trained (mice|rats?|animals|subjects|"
                 r"individuals|men|women|adults|athletes|participants|group|state|muscle|leg)|\btraining (program|"
                 r"protocol|intervention|regimen|sessions?|adaptation|status|load|volume|intensity)|resistance train|"
                 r"strength train|aerobic|treadmill|wheel[- ]running|voluntary running|\brunning\b|swim|cycling|"
                 r"sedentar|athlet|fitness|\bHIIT\b|interval train|detrain|muscle contraction|contractile activity|"
                 r"overload|unloading|disuse|bed rest|inactivity|sprint|\bsports?\b|post-?run|runners?\b|"
                 r"physical (function|performance|capacity)|exerkine|myokine|irisin")
# The paper's own exercise groups, protocol or comparison (contrast_evidence);
# --contrast "" turns the check off.
_EX_WORD = (r"(exercis\w*|train\w*|running|runners?|treadmill|wheels?|swim\w*|cycling|sprint\w*|endurance|"
            r"resistance|aerobic|HIIT|interval|sedentar\w*|athletes?|bouts?|physical(ly)? (activ\w*|inactiv\w*)|"
            r"detrain\w*|unloading|immobili[sz]\w*|bed ?rest|hindlimb (suspension|unloading)|disuse|overload|"
            r"contractions?|electrical stimulation)")
DEFAULT_CONTRAST = "|".join(f"(?:{p})" for p in (
    # "exercise group", "sedentary mice", "trained legs"
    r"\b(exercis\w*|train\w*|running|runners?|treadmill|wheel|swim\w*|sedentar\w*|endurance|resistance|HIIT|"
    r"active|inactive|immobili[sz]ed|unloaded|athletes?) (group|groups|mice|rats?|animals|participants|subjects|"
    r"cohort|arms?|conditions?|legs?|limbs?|volunteers|men|women|individuals|patients|state)\b",
    # "were randomized to ... wheel running", "underwent 8 weeks of training"
    r"(underw\w+|complet\w+|perform\w+|subjected to|assigned (to|into)|randomi[sz]ed (to|into)|allocated (to|into)|"
    r"participated in|enrolled in|given access to|housed with)[^.]{0,80}" + _EX_WORD,
    r"\b(acute|single|one|an?) (bout|session)s? of",
    r"\b(pre|post)[- ]?(exercise|training|intervention|run|bout)",
    r"before and (immediately |\d+ ?h(ours?)? )?after[^.]{0,60}" + _EX_WORD,
    r"(exercise|training|running|treadmill|swim\w*|wheel[- ]running|endurance|resistance|interval)"
    r"([- ]for[- ]\w+\W{0,2})?[- ](protocol|program|programme|intervention|regimen|paradigm|sessions?)",
    # "mice were trained", "participants were exercised"
    r"\b(mice|rats?|animals|participants|subjects|volunteers|patients|men|women) (were|was) "
    r"(trained|exercised|detrained|immobili[sz]ed)",
    r"\b(\d+|one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|sixteen)[- ](week|wk|day|month)s?[- ]"
    r"(of )?([a-z]+\W{1,2}){0,3}(exercise|training|wheel|treadmill|running|endurance|resistance|swim|immobili|unloading|"
    r"suspension|bed ?rest)",
    r"\b(exercised|trained|untrained|detrained) (mice|rats?|animals|group|subjects|participants|individuals|men|"
    r"women|state|muscle|leg)",
    # Associational designs: activity tertiles, athletes vs controls
    r"(tertiles?|quartiles?|levels?|groups?) of (habitual )?physical activity|physical activity "
    r"(tertiles?|quartiles?|levels?|groups?)",
    r"(athletes|trained|active|sedentary) (vs\.?|versus|and|compared (with|to)) (\w+ )?(controls?|untrained|"
    r"sedentary|inactive|non-?athletes)",
))


def screen_file(papers_csv: str, out_dir: str, limit: int = 0, check_design: bool = True,
                routes: set = SUPPORTED_ROUTES, organisms: set | None = None,
                topic: re.Pattern | None = None, contrast: re.Pattern | None = None,
                counts_only: bool = False) -> str:
    os.makedirs(out_dir, exist_ok=True)
    http = Http(cache_dir=os.path.join(out_dir, "_http_cache"))
    with open(papers_csv, encoding="utf-8") as f:
        papers = list(csv.DictReader(f))
    if limit:
        papers = papers[:limit]
    out_csv = os.path.join(out_dir, "screen.csv")
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        for i, row in enumerate(papers, 1):
            try:
                doc = load_full_text(row, os.path.join(out_dir, "papers"), http)
                s = screen_paper(row, doc, http, check_design=check_design, routes=routes,
                                 organisms=organisms, topic=topic, contrast=contrast, counts_only=counts_only)
            except Exception as e:  # one bad paper never stops the batch
                s = Screen(paper_id=row.get("paper_id", ""), title=row.get("title", ""),
                           source=row.get("source", ""), query=row.get("query", ""),
                           verdict="needs_inspection", reason=f"error:{type(e).__name__}: {e}"[:200])
            w.writerow(s.row())
            f.flush()
            print(f"[{i}/{len(papers)}] {s.verdict:16} {s.reason[:40]:40} {s.title[:60]}")
    mark_shared_series(out_csv)
    summary = summarize(out_csv)
    with open(os.path.join(out_dir, "screen_summary.md"), "w", encoding="utf-8") as f:
        f.write(summary)
    return out_csv


def mark_shared_series(screen_csv: str) -> None:
    """Flag analysable GEO series that several papers in the run claim as their own.

    A series is deposited by one paper; when two or more papers claim it, the others
    almost always reuse it and the ownership call missed the reuse wording. Those rows
    become ``needs_inspection`` with reason ``shared_gse`` (the claimants are listed).
    """
    with open(screen_csv, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    claim: dict[str, list] = {}
    for r in rows:
        for g in filter(None, r["own_gse"].split(";")):
            claim.setdefault(g, []).append(r["paper_id"])
    changed = False
    for r in rows:
        shared = [g for g in filter(None, r["analyzable_gse"].split(";")) if len(claim[g]) > 1]
        if shared and r["verdict"] == "analyzable":
            r["verdict"] = "needs_inspection"
            r["reason"] = "shared_gse:" + ",".join(f"{g}({len(claim[g])} papers)" for g in shared)
            changed = True
    if changed:
        with open(screen_csv, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=FIELDS)
            w.writeheader()
            w.writerows(rows)


def summarize(screen_csv: str) -> str:
    with open(screen_csv, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    n = len(rows)
    verdicts = Counter(r["verdict"] for r in rows)
    reasons = Counter(r["reason"].split(":")[0] for r in rows if r["verdict"] == "excluded")
    types = Counter(t for r in rows for t in r["data_types"].split(";") if t)
    lines = [f"# Paper screening summary ({n} papers)", "", "| Verdict | Papers |", "|---|---:|"]
    lines += [f"| {k} | {v} |" for k, v in verdicts.most_common()]
    lines += ["", "## Why papers were excluded (first failing layer)", "", "| Reason | Papers |", "|---|---:|"]
    lines += [f"| {k} | {v} |" for k in REASONS for v in [reasons.get(k, 0)] if v]
    lines += [f"| {k} | {v} |" for k, v in reasons.items() if k not in REASONS]
    inspect = Counter(r["reason"].split(":")[0] for r in rows if r["verdict"] == "needs_inspection")
    if inspect:
        lines += ["", "## Why papers need inspection", "", "| Reason | Papers |", "|---|---:|"]
        lines += [f"| {k} | {v} |" for k, v in inspect.most_common()]
    in_tar = sum(1 for r in rows if r["verdict"] == "analyzable"
                 and any((x.get("files") or {}).get("status") == "raw_counts_in_tar"
                         for x in json.loads(r.get("series_json") or "[]") if x["gse"] in r["analyzable_gse"]))
    if in_tar:
        lines += ["", f"{in_tar} analyzable papers have their raw counts only inside a _RAW.tar."]
    sources = sorted({r.get("source", "") for r in rows} - {""})
    if len(sources) > 1:
        lines += ["", "## Verdicts by source", "", "| Source | Papers | analyzable | needs_inspection |",
                  "|---|---:|---:|---:|"]
        for src in sources:
            sub = [r for r in rows if r.get("source") == src]
            lines.append(f"| {src} | {len(sub)} | {sum(r['verdict'] == 'analyzable' for r in sub)} | "
                         f"{sum(r['verdict'] == 'needs_inspection' for r in sub)} |")
    lines += ["", "## Data types of the papers' own GEO series", "", "| Data type | Papers |", "|---|---:|"]
    lines += [f"| {k} | {v} |" for k, v in types.most_common()]
    ok = [r for r in rows if r["verdict"] == "analyzable"]
    if ok:
        routes = Counter(r.get("route", "") for r in ok)
        lines += ["", "## Analysable papers by route", "", "| Route | Papers |", "|---|---:|"]
        lines += [f"| {k or '?'} | {v} |" for k, v in routes.most_common()]
        flags = Counter(f.split(":")[0] if not f.startswith("sex:") else f
                        for r in ok for f in filter(None, r.get("flags", "").split(";")))
        lines += ["", "## Checks for the analysable papers", "", "| Flag | Papers |", "|---|---:|"]
        lines += [f"| {k} | {v} |" for k, v in sorted(flags.items())]
    lines += ["", "## What each data type needs", "", "| Data type | Needed |", "|---|---|"]
    lines += [f"| {k} | {v} |" for k, v in REQUIRED_DATA.items()]
    return "\n".join(lines) + "\n"


def _main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="python -m tools.paper_screen", description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("collect", help="search Semantic Scholar + PubMed, write papers.csv")
    p.add_argument("query", nargs="+", help="one or more queries; results are merged and deduplicated")
    p.add_argument("--out", required=True)
    p.add_argument("--max", type=int, default=200, help="total papers across all queries and sources")
    p.add_argument("--end-year", type=int, default=0)
    p.add_argument("--start-year", type=int, default=0)
    p.add_argument("--shares", default="",
                   help="per-source share of --max, e.g. geo=0.7,s2=0.15,pubmed=0.15")
    p.add_argument("--exclude", default="", help="papers.csv of an earlier run; its papers are skipped")
    p.add_argument("--sources", default="s2,pubmed",
                   help="comma-separated: s2, pubmed, geo (papers linked from GEO bulk RNA-seq series)")
    p = sub.add_parser("screen", help="screen a papers.csv, write screen.csv + screen_summary.md")
    p.add_argument("papers_csv")
    p.add_argument("--out", required=True)
    p.add_argument("--limit", type=int, default=0)
    p.add_argument("--no-design", action="store_true", help="skip the series-matrix design layer")
    p.add_argument("--counts-only", action="store_true",
                   help="exclude series with only TPM/FPKM/log matrices instead of routing them to limma")
    p.add_argument("--routes", default="rnaseq", help="comma-separated supported data types")
    p.add_argument("--organisms", default=",".join(sorted(DEFAULT_ORGANISMS)),
                   help='comma-separated species kept (default human, mouse, rat); "any" keeps all')
    p.add_argument("--topic", default=DEFAULT_TOPIC,
                   help='regex the title or abstract must match (default: exercise terms); "" turns it off')
    p.add_argument("--contrast", default=DEFAULT_CONTRAST,
                   help="regex for the paper's own exercise groups or protocol in its abstract, methods "
                        'or results (default: exercise designs); "" turns it off')
    p = sub.add_parser("summary")
    p.add_argument("screen_csv")
    a = ap.parse_args(argv)
    if a.cmd == "collect":
        os.makedirs(a.out, exist_ok=True)
        sources = [x for x in a.sources.split(",") if x]
        bad = set(sources) - set(SOURCES)
        if bad:
            ap.error(f"unknown source(s) {sorted(bad)}; choose from {', '.join(SOURCES)}")
        shares = {k: float(v) for k, v in (x.split("=") for x in a.shares.split(",") if x)} or None
        exclude = list(csv.DictReader(open(a.exclude, encoding="utf-8"))) if a.exclude else None
        rows = collect_papers(a.query, Http(cache_dir=os.path.join(a.out, "_http_cache")), a.max,
                              a.end_year, sources, start_year=a.start_year, shares=shares, exclude=exclude)
        path = os.path.join(a.out, "papers.csv")
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else ["paper_id"])
            w.writeheader()
            w.writerows(rows)
        print(f"{len(rows)} papers -> {path}")
    elif a.cmd == "screen":
        organisms = None if a.organisms.strip().lower() == "any" else \
            {o.strip().lower() for o in a.organisms.split(",") if o.strip()}
        topic = re.compile(a.topic, re.I) if a.topic else None
        out = screen_file(a.papers_csv, a.out, a.limit, not a.no_design, set(a.routes.split(",")),
                          organisms, topic, re.compile(a.contrast, re.I) if a.contrast else None,
                          a.counts_only)
        print(open(os.path.join(a.out, "screen_summary.md"), encoding="utf-8").read())
        print(out)
    else:
        print(summarize(a.screen_csv))
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except Exception:
            pass
    sys.exit(_main())
