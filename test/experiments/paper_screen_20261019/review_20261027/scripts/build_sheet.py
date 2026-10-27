"""Build the manual review sheet (xlsx) for the 58 sampled papers."""
import csv, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); sys.path.insert(0, os.getcwd())
from judgments import J

RUN = "test/experiments/paper_screen_20261019"
REV = os.path.join(RUN, "review_20261027")
papers = {r["paper_id"]: r for r in csv.DictReader(open(os.path.join(RUN, "papers.csv"), encoding="utf-8"))}
screen = {r["paper_id"]: r for r in csv.DictReader(open(os.path.join(REV, "screen.csv"), encoding="utf-8"))}
sample = list(csv.DictReader(open(os.path.join(REV, "sample.csv"), encoding="utf-8")))

LAYER = {
    "no_full_text": "1 无全文", "off_topic": "2 非运动主题", "no_exercise_contrast": "3 无运动分组",
    "no_own_deposit": "4 无自有数据", "no_own_deposit:reuses_public_gse": "4 无自有数据：重用公共 GSE",
    "no_own_deposit:linked_to_other_paper": "4 无自有数据：GSE 属其他文章", "data_on_request": "4 无自有数据：需索取",
    "non_geo_repository": "4 无自有数据：非 GEO", "unsupported_data_type:microarray": "5 非 bulk RNA-seq：芯片",
    "unsupported_data_type:scrnaseq": "5 非 bulk RNA-seq：单细胞", "unsupported_data_type:other": "5 非 bulk RNA-seq：其他",
    "unsupported_organism": "6 物种", "no_processed_matrix": "7 无计数：无矩阵", "de_tables_only": "7 无计数：只有差异表",
    "normalized_only": "7 无计数：只有 FPKM/TPM", "single_group": "8 分组：只有一组", "group_too_small": "8 分组：组太小",
    "matrix_unlabelled_matrix": "9 待看：矩阵未标注", "shared_gse": "9 待看：多篇共用 GSE", "matrix_raw_tar": "9 待看：RAW.tar",
    "matrix_decimal_unexplained": "9 待看：小数计数", "error": "9 待看：出错", "analyzable": "10 可分析（对照）",
}

import re, shutil
from tools import paper_screen as ps
from tools import paper_structure as pst
FT = os.path.join(REV, "fulltext"); os.makedirs(FT, exist_ok=True)
CONTRAST = re.compile(ps.DEFAULT_CONTRAST, re.I)
ACC = re.compile(r"\b(GSE\d+|PRJ[NED][A-Z]?\d+|SR[PAR]\d+|E-[A-Z]{4}-\d+|PXD\d+|phs\d+)")

def stem_of(p):
    return re.sub(r"[^A-Za-z0-9._-]", "_", p.get("pmcid") or p.get("paper_id") or "paper")

def fulltext(p, tag):
    """Copy the cached full text (and PDF) next to the sheet; return the relative path or ''."""
    st = stem_of(p); src = os.path.join(RUN, "papers", st)
    if not os.path.exists(src + ".paper.md"):
        return ""
    name = f"{tag}_{p['pmid'] or st[:12]}"
    shutil.copy(src + ".paper.md", os.path.join(FT, name + ".md"))
    if os.path.exists(src + ".pdf"):
        shutil.copy(src + ".pdf", os.path.join(FT, name + ".pdf"))
    return f"fulltext/{name}.md"

def excerpt(p, limit=1500):
    """Data statement, sentences naming accessions, and the exercise-design phrases."""
    st = stem_of(p); path = os.path.join(RUN, "papers", st + ".paper.json")
    if not os.path.exists(path):
        return "（没有取到全文）"
    blocks = pst.load(path).get("blocks", [])
    out = []
    for b in blocks:
        if b.get("section_class") == "data_availability" and b["text"].strip():
            out.append("【数据声明】" + b["text"][:500])
    for b in blocks:
        if b.get("section_class") == "data_availability": continue
        for sent in re.split(r"(?<=[.;])\s+", b["text"]):
            if ACC.search(sent) and len(out) < 6:
                out.append(f"【{b.get('section_class')}】" + sent[:300])
    hits = ps.contrast_evidence({"blocks": blocks}, CONTRAST, limit=2)
    by_id = {b.get("id"): b for b in blocks}
    for h in hits:
        bid, phrase = h.split(": ", 1)
        t = by_id.get(bid, {}).get("text", ""); i = t.find(phrase)
        out.append("【运动设计】…" + t[max(0, i - 120): i + 200] + "…")
    if not hits:
        out.append("【运动设计】方法/结果里没找到运动分组或方案的句子")
    return "\n".join(dict.fromkeys(out))[:limit]

def why(s):
    r, own, cited, other = s["reason"], s["own_gse"], s["cited_gse"], s["other_repos"]
    series = json.loads(s["series_json"] or "[]")
    if r == "no_full_text": return "PMC 和 PDF 都没取到全文"
    if r == "off_topic": return "标题和摘要里没有运动相关词"
    if r == "no_exercise_contrast": return "摘要、方法、结果里找不到本文自己的运动分组或方案"
    if r == "no_own_deposit": return "正文没有任何数据库编号"
    if r == "no_own_deposit:reuses_public_gse": return f"正文提到的 GSE 被判为引用：{cited[:60]}"
    if r == "data_on_request": return "数据声明写“向作者索取”"
    if r.startswith("non_geo_repository"): return f"只给了非 GEO 编号：{other[:60]}"
    if r.startswith("error"): return r[:80]
    parts = []
    for se in series:
        f, d = se.get("files") or {}, se.get("design") or {}
        bits = [se["reason"] or "analyzable", se["data_type"], se["taxon"]]
        if f.get("status"): bits.append("文件:" + f["status"])
        if d.get("groups"): bits.append("分组:" + ",".join(f"{k or '空'}={v}" for k, v in list(d["groups"].items())[:5]))
        parts.append(f"{se['gse']}（" + "；".join(b for b in bits if b) + "）")
    return "\n".join(parts)[:400] or r

wb = Workbook(); ws = wb.active; ws.title = "复核"
head = ["#", "筛选层", "标题", "年份", "PMID", "相关 GSE", "筛选器的理由", "我的判断", "我认为能跑吗", "我的说明",
        "原文摘录", "全文（本地）", "全文（PMC）", "DOI", "GEO 链接", "你的结论：能跑吗", "你的备注"]
ws.append(head)
for i, row in enumerate(sample, 1):
    p, s = papers[row["paper_id"]], screen[row["paper_id"]]
    ok, run, gse, note = J[i]
    gses = [g for g in (gse or s["own_gse"]).split(";") if g.startswith("GSE")]
    local = fulltext(p, f"{i:02d}")
    pmc = f"https://pmc.ncbi.nlm.nih.gov/articles/{p['pmcid']}/" if p["pmcid"] else ""
    doi = f"https://doi.org/{p['doi']}" if p["doi"] else ""
    geo = f"https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc={gses[0]}" if gses else ""
    ws.append([i, LAYER[row["layer"]], p["title"], p["year"], p["pmid"], gse or s["own_gse"], why(s), ok, run, note,
               excerpt(p), local or "（无）", pmc, doi, geo, "", ""])
    for col, url in ((12, local), (13, pmc), (14, doi), (15, geo)):
        if url:
            c = ws.cell(row=i + 1, column=col); c.hyperlink = url; c.style = "Hyperlink"
widths = [4, 22, 40, 6, 10, 16, 40, 10, 12, 50, 80, 18, 18, 18, 18, 14, 30]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w
for r in ws.iter_rows(min_row=2):
    for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
for c in ws[1]:
    c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
yellow = PatternFill("solid", fgColor="FFF2CC")
red = PatternFill("solid", fgColor="F8CBAD")
for r in range(2, len(sample) + 2):
    ws.cell(row=r, column=16).fill = yellow; ws.cell(row=r, column=17).fill = yellow
    if ws.cell(row=r, column=8).value in ("错", "存疑"): ws.cell(row=r, column=8).fill = red
dv = DataValidation(type="list", formula1='"能,不能,也许,看不出"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"P2:P{len(sample) + 1}")
ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:Q{len(sample) + 1}"
# Sheet 2: papers GEO links to its own RNA-seq series, dropped before the data layers
ws2 = wb.create_sheet("GEO反查可补回")
ws2.append(["#", "原筛选层", "标题", "标题有运动词", "PMID", "GSE", "按数据层判", "文件", "分组", "原文摘录",
            "全文（本地）", "全文（PMC）", "DOI", "GEO 链接", "你的结论：能跑吗", "你的备注"])
wi = [r for r in csv.DictReader(open(os.path.join(REV, "whatif_geolink.csv"), encoding="utf-8")) if r["verdict"] != "excluded"]
seen = {}
for r in wi:
    seen.setdefault(r["paper_id"], []).append(r)
k = 0
for pid, rs in sorted(seen.items(), key=lambda x: (x[1][0]["layer"], x[0])):
    k += 1
    r = rs[0]; gse = ";".join(x["gse"] for x in rs)
    url = f"https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc={rs[0]['gse']}"
    p = papers[pid]; local = fulltext(p, f"B{k:02d}")
    pmc = f"https://pmc.ncbi.nlm.nih.gov/articles/{p['pmcid']}/" if p["pmcid"] else ""
    doi = f"https://doi.org/{p['doi']}" if p["doi"] else ""
    ws2.append([k, LAYER.get(r["layer"], r["layer"]), r["title"], "是" if r["title_on_topic"] == "True" else "否",
                p["pmid"], gse, ";".join((x["verdict"] + (":" + x["reason"] if x["reason"] else "")) for x in rs),
                ";".join(x["files"] for x in rs), r["groups"], excerpt(p), local or "（无）", pmc, doi, url, "", ""])
    for col, u in ((11, local), (12, pmc), (13, doi), (14, url)):
        if u:
            c = ws2.cell(row=k + 1, column=col); c.hyperlink = u; c.style = "Hyperlink"
for j, w in enumerate([4, 22, 45, 8, 10, 14, 22, 14, 36, 70, 18, 18, 18, 18, 14, 30], 1):
    ws2.column_dimensions[ws2.cell(row=1, column=j).column_letter].width = w
for row in ws2.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
for c in ws2[1]:
    c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
for r in range(2, k + 2):
    ws2.cell(row=r, column=15).fill = yellow; ws2.cell(row=r, column=16).fill = yellow
dv2 = DataValidation(type="list", formula1='"能,不能,也许,看不出"', allow_blank=True)
ws2.add_data_validation(dv2); dv2.add(f"O2:O{k + 1}")
ws2.freeze_panes = "D2"
out = os.path.join(REV, "review_sheet.xlsx"); wb.save(out); print(out)
